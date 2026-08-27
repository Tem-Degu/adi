# -*- coding: utf-8 -*-
"""Shared design system for the StreetAI decks.

Both `build_deck.py` (full deck) and `build_deck_investor.py` (investor cut)
import this, so the two stay visually identical and only the content differs.

Copy rules, from user feedback, partly enforced by `check_deck.py`:

REGISTER: write as a technical report, not as marketing. State the claim and
stop. No enthusiasm, no persuasion in the prose itself. The numbers persuade.

  - No em dashes or en dashes.
  - No "X, not Y" construction.
  - No closing aphorism on slides. A slide ends when its content ends.
  - No second person and no imperatives. Do not address the reader.
  - No metaphors, similes or rhetorical questions. "The machine is a heater
    with a pump" is the kind of line to delete.
  - No sentence fragments used for emphasis.
  - Complete declarative sentences. Subject, verb, object.
  - Every quantity carries a figure and, where external, a source.
  - Headings name their content. They are labels, not slogans.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---------------------------------------------------------------- palette
BG      = RGBColor(0x0E, 0x0F, 0x12)
CARD    = RGBColor(0x1A, 0x1C, 0x22)
CARD_HI = RGBColor(0x24, 0x27, 0x2F)
TEXT    = RGBColor(0xF2, 0xF0, 0xEC)
MUTED   = RGBColor(0x93, 0x96, 0xA0)
ACCENT  = RGBColor(0xE0, 0x8A, 0x3C)
TEAL    = RGBColor(0x6F, 0xA8, 0xA0)
FONT    = "Segoe UI"

W, H = Inches(13.333), Inches(7.5)
M    = Inches(0.85)
CW   = W - 2 * M

# Standard column geometries.
CW2      = Inches(5.35)
CW3, G3  = Inches(3.66), Inches(0.29)
CW4, G4  = Inches(2.72), Inches(0.22)


class Deck:
    """A 16:9 deck with the StreetAI look."""

    def __init__(self, footer_text="StreetAI  ·  Artificial Domain Intelligence"):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H
        self.blank = self.prs.slide_layouts[6]
        self.footer_text = footer_text
        self.count = 0

    # ------------------------------------------------------------ slides
    def slide(self, footer=True, number=True):
        s = self.prs.slides.add_slide(self.blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = BG
        self.count += 1
        if footer:
            self.tb(s, self.footer_text, M, H - Inches(0.62), Inches(6),
                    Inches(0.3), 10, MUTED, space_after=0)
        if number:
            self.tb(s, f"{self.count:02d}", W - M - Inches(1), H - Inches(0.62),
                    Inches(1), Inches(0.3), 10, MUTED, align=PP_ALIGN.RIGHT,
                    space_after=0)
        return s

    # ------------------------------------------------------------- text
    def tb(self, s, text, x, y, w, h, size, color, bold=False,
           align=PP_ALIGN.LEFT, line=1.15, space_after=6,
           anchor=MSO_ANCHOR.TOP, caps=False):
        box = s.shapes.add_textbox(x, y, w, h)
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = anchor
        p = tf.paragraphs[0]
        p.alignment = align
        p.line_spacing = line
        p.space_after = Pt(space_after)
        r = p.add_run()
        r.text = text.upper() if caps else text
        r.font.size, r.font.bold, r.font.name = Pt(size), bold, FONT
        r.font.color.rgb = color
        return box

    @staticmethod
    def _hang(p, inches):
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(int(Inches(inches))))
        pPr.set("indent", str(-int(Inches(inches))))

    def bullets(self, s, items, x, y, w, h, size=17, color=TEXT, gap=16,
                line=1.3, marker=True, marker_color=None, hang=0.26):
        """items: str, or (bold_lead, rest) tuple."""
        box = s.shapes.add_textbox(x, y, w, h)
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        mc = marker_color or ACCENT
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.line_spacing = line
            p.space_after = Pt(gap)
            if marker:
                self._hang(p, hang)
                r = p.add_run()
                r.text = "•\t"
                r.font.size, r.font.name = Pt(size), FONT
                r.font.color.rgb = mc
            if isinstance(item, tuple):
                lead, rest = item
                r = p.add_run()
                r.text = lead
                r.font.size, r.font.bold, r.font.name = Pt(size), True, FONT
                r.font.color.rgb = TEXT
                r2 = p.add_run()
                r2.text = rest
                r2.font.size, r2.font.name = Pt(size), FONT
                r2.font.color.rgb = MUTED
            else:
                r = p.add_run()
                r.text = item
                r.font.size, r.font.name = Pt(size), FONT
                r.font.color.rgb = color
        return box

    # ------------------------------------------------------------ shapes
    def rect(self, s, x, y, w, h, fill=CARD):
        sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        sh.adjustments[0] = 0.045
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
        sh.line.fill.background()
        sh.shadow.inherit = False
        sh.text_frame.text = ""
        return sh

    def bar(self, s, x, y, w=Inches(0.9), h=Inches(0.05), color=ACCENT):
        sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
        sh.fill.solid()
        sh.fill.fore_color.rgb = color
        sh.line.fill.background()
        sh.shadow.inherit = False
        return sh

    # ---------------------------------------------------------- compounds
    def heading(self, s, kicker, title):
        self.tb(s, kicker, M, Inches(0.72), CW, Inches(0.3), 11.5, ACCENT,
                bold=True, caps=True)
        self.tb(s, title, M, Inches(1.14), CW, Inches(1.1), 34, TEXT,
                bold=True, line=1.08)

    def card(self, s, x, y, w, h, title, body, accent=ACCENT, ts=17, bs=13):
        self.rect(s, x, y, w, h)
        self.bar(s, x + Inches(0.3), y + Inches(0.34), Inches(0.5),
                 Inches(0.045), accent)
        self.tb(s, title, x + Inches(0.3), y + Inches(0.58), w - Inches(0.6),
                Inches(0.5), ts, TEXT, bold=True, line=1.15)
        self.tb(s, body, x + Inches(0.3), y + Inches(1.12), w - Inches(0.6),
                h - Inches(1.4), bs, MUTED, line=1.4)

    def panel(self, s, x, y, w, h, title, body, accent=ACCENT, ts=15, bs=12.5):
        """A shorter card. Same look, less vertical reserve than `card`."""
        self.rect(s, x, y, w, h)
        self.bar(s, x + Inches(0.28), y + Inches(0.26), Inches(0.4),
                 Inches(0.04), accent)
        self.tb(s, title, x + Inches(0.28), y + Inches(0.44), w - Inches(0.56),
                Inches(0.35), ts, TEXT, bold=True, line=1.15)
        self.tb(s, body, x + Inches(0.28), y + Inches(0.82), w - Inches(0.56),
                h - Inches(1.05), bs, MUTED, line=1.38)

    def picture(self, s, path, x, y, w, h, frame=True):
        """Place an image, optionally with a thin frame so it reads as a panel
        rather than a cut-out on the dark background."""
        pic = s.shapes.add_picture(str(path), x, y, width=w, height=h)
        if frame:
            pic.line.color.rgb = CARD_HI
            pic.line.width = Pt(1)
        return pic

    def stat(self, s, x, y, w, h, figure, label, note=None, accent=ACCENT):
        """A big number with a label under it."""
        self.rect(s, x, y, w, h)
        self.tb(s, figure, x + Inches(0.3), y + Inches(0.34), w - Inches(0.6),
                Inches(0.8), 34, accent, bold=True, line=1.0)
        self.tb(s, label, x + Inches(0.3), y + Inches(1.16), w - Inches(0.6),
                Inches(0.6), 14, TEXT, bold=True, line=1.2)
        if note:
            self.tb(s, note, x + Inches(0.3), y + Inches(1.82),
                    w - Inches(0.6), h - Inches(2.1), 12, MUTED, line=1.35)

    def row(self, s, y, height, cells, fill=CARD, width=Inches(11.0)):
        """A full-width band. cells: list of (x_offset_in, width_in, text,
        size, color, bold)."""
        self.rect(s, M, y, width, height, fill)
        for x_off, w_in, text, size, color, bold in cells:
            self.tb(s, text, M + Inches(x_off), y + Inches(0.24), Inches(w_in),
                    height - Inches(0.4), size, color, bold=bold, line=1.3)

    def save(self, path):
        self.prs.save(path)
        print(f"saved: {path} | slides: {self.count}")
